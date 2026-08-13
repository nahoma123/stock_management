import base64
import csv
import io
import re

from odoo import _, api, fields, models
from odoo.exceptions import UserError, ValidationError


SUPPORTED_EXTENSIONS = ('.csv', '.xlsx')
BASE_COLUMNS = {
    'machinecode': 'machine_code',
    'modelcode': 'model_code',
    'description': 'description',
    'priceetb': 'price_etb',
    'priceusd': 'price_usd',
    'sales': 'sales',
    'total': 'total',
}


def _clean_header(value):
    return re.sub(r'[^a-z0-9]', '', str(value or '').strip().lower())


def _text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _number(value, label, required=False):
    raw = _text(value).replace(',', '').replace('$', '').replace('ETB', '').strip()
    if not raw:
        if required:
            raise ValidationError(_('%s is required.') % label)
        return 0.0
    try:
        return float(raw)
    except ValueError as error:
        raise ValidationError(_('%(label)s must be a number, got "%(value)s".', label=label, value=value)) from error


class InitialDataImportWizard(models.TransientModel):
    _name = 'initial.data.import.wizard'
    _description = 'Initial Product Data Import'

    file = fields.Binary(required=True, attachment=False)
    filename = fields.Char(required=True)
    state = fields.Selection([('upload', 'Upload'), ('preview', 'Preview'), ('done', 'Done')], default='upload', required=True)
    line_ids = fields.One2many('initial.data.import.line', 'wizard_id')
    imported_count = fields.Integer(readonly=True)
    warning_message = fields.Text(readonly=True)

    def action_preview(self):
        self.ensure_one()
        self._validate_file()
        rows, image_map = self._read_file()
        columns, data_start = self._detect_columns(rows)
        commands = [fields.Command.clear()]
        for row_number, row in enumerate(rows[data_start:], start=data_start + 1):
            if not any(_text(cell) for cell in row):
                continue
            values = self._prepare_line(row, row_number, columns, image_map)
            commands.append(fields.Command.create(values))
        if len(commands) == 1:
            raise UserError(_('The file contains no product rows.'))
        self.write({'line_ids': commands, 'state': 'preview', 'warning_message': False})
        return self._reload()

    def action_import(self):
        self.ensure_one()
        if self.state != 'preview' or not self.line_ids:
            raise UserError(_('Preview the file before importing it.'))
        invalid_lines = self.line_ids.filtered(lambda line: line.status == 'error')
        if invalid_lines:
            raise UserError(_('Fix or remove rows with errors before importing.'))

        stock_locations = {}
        pricelists = {
            'ETB': self._get_pricelist('ETB'),
            'USD': self._get_pricelist('USD'),
        }
        imported = 0
        for line in self.line_ids.filtered(lambda item: item.include):
            product = self._upsert_product(line)
            self._set_pricelist_price(pricelists['ETB'], product, line.price_etb)
            self._set_pricelist_price(pricelists['USD'], product, line.price_usd)
            for location_name, quantity in line.stock_values().items():
                if not quantity:
                    continue
                location = stock_locations.setdefault(location_name, self._get_stock_location(location_name))
                self._set_opening_stock(product.product_variant_id, location, quantity)
            imported += 1

        self.write({
            'state': 'done',
            'imported_count': imported,
            'warning_message': _('Imported %s product rows successfully.') % imported,
        })
        return self._reload()

    def action_reset(self):
        self.ensure_one()
        self.write({'state': 'upload', 'line_ids': [fields.Command.clear()], 'warning_message': False})
        return self._reload()

    def _validate_file(self):
        filename = (self.filename or '').lower()
        if not filename.endswith(SUPPORTED_EXTENSIONS):
            raise UserError(_('Upload a CSV or XLSX file.'))
        if not self.file:
            raise UserError(_('Select a file to import.'))

    def _read_file(self):
        content = base64.b64decode(self.file)
        if self.filename.lower().endswith('.csv'):
            return self._read_csv(content), {}
        return self._read_xlsx(content)

    def _read_csv(self, content):
        text = content.decode('utf-8-sig', errors='replace')
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        return [list(row) for row in csv.reader(io.StringIO(text), dialect)]

    def _read_xlsx(self, content):
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise UserError(_('XLSX support is not installed on the server.')) from error
        try:
            sheet = load_workbook(io.BytesIO(content), data_only=True).active
        except Exception as error:
            raise UserError(_('The XLSX file could not be read: %s') % error) from error
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        images = {}
        for image in getattr(sheet, '_images', []):
            try:
                images[(image.anchor._from.row, image.anchor._from.col)] = base64.b64encode(image._data())
            except (AttributeError, ValueError):
                continue
        return rows, images

    def _detect_columns(self, rows):
        for header_index, row in enumerate(rows[:10]):
            normalized = [_clean_header(cell) for cell in row]
            if 'machinecode' not in normalized:
                continue
            columns = {}
            picture_number = 0
            stock_number = 0
            for index, header in enumerate(normalized):
                upper_header = ''
                if header_index:
                    previous_row = rows[header_index - 1]
                    if index < len(previous_row):
                        upper_header = _clean_header(previous_row[index])
                effective_header = header or upper_header
                if header in BASE_COLUMNS:
                    columns[BASE_COLUMNS[header]] = index
                elif header.startswith('picture'):
                    picture_number += 1
                    columns[f'picture_{picture_number}'] = index
                elif effective_header.startswith('stock'):
                    stock_number += 1
                    suffix = re.sub(r'^stock', '', effective_header) or str(stock_number)
                    columns[f'stock_{suffix}'] = index
            if 'description' not in columns:
                raise UserError(_('The header must contain a Description column.'))
            return columns, header_index + 1
        raise UserError(_('Could not find the header row. It must contain Machine Code and Description.'))

    def _prepare_line(self, row, row_number, columns, image_map):
        def cell(name):
            index = columns.get(name)
            return row[index] if index is not None and index < len(row) else None

        errors = []
        warnings = []
        description = _text(cell('description'))
        machine_code = _text(cell('machine_code'))
        if not description:
            errors.append(_('Description is required.'))
        if not machine_code:
            warnings.append(_('Machine Code is empty; this row cannot be matched on a later re-import.'))
        try:
            price_etb = _number(cell('price_etb'), _('PRICE ETB'))
            price_usd = _number(cell('price_usd'), _('PRICE USD'))
        except ValidationError as error:
            errors.append(str(error))
            price_etb = price_usd = 0.0

        stocks = {}
        for key, index in columns.items():
            if not key.startswith('stock_'):
                continue
            try:
                stocks[key] = _number(row[index] if index < len(row) else None, key.upper())
                if stocks[key] < 0:
                    errors.append(_('%s cannot be negative.') % key.upper())
            except ValidationError as error:
                errors.append(str(error))
                stocks[key] = 0.0

        picture = False
        for key, index in columns.items():
            if key.startswith('picture_') and (row_number - 1, index) in image_map:
                picture = image_map[(row_number - 1, index)]
                break
        if not picture and any(key.startswith('picture_') and _text(cell(key)) for key in columns):
            warnings.append(_('Only pictures embedded in XLSX cells are imported; text paths and CSV picture values are ignored.'))

        return {
            'row_number': row_number,
            'include': not errors,
            'machine_code': machine_code,
            'model_code': _text(cell('model_code')),
            'description': description,
            'price_etb': price_etb,
            'price_usd': price_usd,
            'stock_1': stocks.get('stock_1', 0.0),
            'stock_2': stocks.get('stock_2', 0.0),
            'stock_3': stocks.get('stock_3', 0.0),
            'image_1920': picture,
            'status': 'error' if errors else 'ready',
            'message': '\n'.join(errors + warnings),
        }

    def _upsert_product(self, line):
        Product = self.env['product.template'].with_context(active_test=False)
        product = Product.search([('default_code', '=', line.machine_code)], limit=1) if line.machine_code else Product.browse()
        values = {
            'name': line.description,
            'default_code': line.machine_code or False,
            'model_code': line.model_code or False,
            'list_price': line.price_etb,
            'is_storable': True,
            'active': True,
        }
        if line.image_1920:
            values['image_1920'] = line.image_1920
        if product:
            product.write(values)
        else:
            product = Product.create(values)
        return product

    def _get_pricelist(self, currency_code):
        Currency = self.env['res.currency'].with_context(active_test=False)
        currency = Currency.search([('name', '=', currency_code)], limit=1)
        if not currency and currency_code == 'ETB':
            currency = Currency.create({
                'name': 'ETB',
                'symbol': 'Br',
                'rounding': 0.01,
                'position': 'after',
                'active': True,
            })
        if not currency:
            raise UserError(_('Currency %s is not configured in Odoo.') % currency_code)
        currency.active = True
        pricelist = self.env['product.pricelist'].search([
            ('currency_id', '=', currency.id),
            ('company_id', 'in', [False, self.env.company.id]),
        ], limit=1)
        return pricelist or self.env['product.pricelist'].create({
            'name': _('%s Sales Prices') % currency_code,
            'currency_id': currency.id,
            'company_id': self.env.company.id,
        })

    def _set_pricelist_price(self, pricelist, product, price):
        if not price:
            return
        Item = self.env['product.pricelist.item']
        item = Item.search([('pricelist_id', '=', pricelist.id), ('product_tmpl_id', '=', product.id)], limit=1)
        values = {
            'pricelist_id': pricelist.id,
            'applied_on': '1_product',
            'product_tmpl_id': product.id,
            'compute_price': 'fixed',
            'fixed_price': price,
        }
        item.write(values) if item else Item.create(values)

    def _get_stock_location(self, name):
        warehouse = self.env['stock.warehouse'].search([('company_id', '=', self.env.company.id)], limit=1)
        if not warehouse:
            raise UserError(_('No warehouse is configured for this company.'))
        location = self.env['stock.location'].search([
            ('name', '=', name),
            ('location_id', '=', warehouse.lot_stock_id.id),
            ('company_id', 'in', [False, self.env.company.id]),
        ], limit=1)
        return location or self.env['stock.location'].create({
            'name': name,
            'location_id': warehouse.lot_stock_id.id,
            'usage': 'internal',
            'company_id': self.env.company.id,
        })

    def _set_opening_stock(self, product, location, target_quantity):
        Quant = self.env['stock.quant']
        current = Quant._get_available_quantity(product, location, allow_negative=True)
        difference = target_quantity - current
        if difference:
            Quant._update_available_quantity(product, location, difference)

    def _reload(self):
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }


class InitialDataImportLine(models.TransientModel):
    _name = 'initial.data.import.line'
    _description = 'Initial Product Data Import Preview Line'
    _order = 'row_number'

    wizard_id = fields.Many2one('initial.data.import.wizard', required=True, ondelete='cascade')
    row_number = fields.Integer(readonly=True)
    include = fields.Boolean(default=True)
    machine_code = fields.Char()
    model_code = fields.Char()
    description = fields.Char()
    price_etb = fields.Float()
    price_usd = fields.Float()
    stock_1 = fields.Float(string='STOCK-1')
    stock_2 = fields.Float(string='STOCK-2')
    stock_3 = fields.Float(string='STOCK-3')
    image_1920 = fields.Binary(string='Picture', attachment=False)
    status = fields.Selection([('ready', 'Ready'), ('error', 'Error')], readonly=True)
    message = fields.Text(readonly=True)

    def stock_values(self):
        self.ensure_one()
        return {'STOCK-1': self.stock_1, 'STOCK-2': self.stock_2, 'STOCK-3': self.stock_3}
